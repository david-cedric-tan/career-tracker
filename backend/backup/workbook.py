"""Render an archive to .xlsx and read it back (FR-EXPORT-02).

The spreadsheet is generated from `SHEETS`, and parsed with the same table, so
a workbook this app produced can always be imported by it — the round trip is
structural rather than something that has to be kept in sync by hand.
"""

import json
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from rest_framework import serializers

from .archive import ARCHIVE_VERSION, LIST_SEPARATOR, SHEETS

META_SHEET = "_meta"
HEADER_FILL = PatternFill("solid", fgColor="1F2A37")
HEADER_FONT = Font(color="FFFFFF", bold=True)

# Columns whose values are lists of names, flattened to a delimited string so a
# spreadsheet cell stays human-readable and still parses back.
LIST_COLUMNS = {
    "target_companies",
    "target_roles",
    "listings",
    "companies",
    "applications",
    "photo_captions",
    "industries",
}
# …of which these hold ids, not names. A spreadsheet cell is text, so without
# coercing them back to ints the restore's id lookups silently find nothing.
ID_LIST_COLUMNS = {"applications"}
# Columns carrying structured JSON that must survive the round trip verbatim.
JSON_COLUMNS = {"changes"}
BOOL_COLUMNS = {"is_active", "is_preferred"}
INT_COLUMNS = {"id", "application"}


def _to_cell(column, value):
    if value is None:
        return ""
    if column in JSON_COLUMNS:
        return json.dumps(value)
    if column in LIST_COLUMNS:
        return LIST_SEPARATOR.join(str(item) for item in value)
    if isinstance(value, bool):
        return "yes" if value else "no"
    return value


def _from_cell(column, value):
    text = "" if value is None else str(value).strip()

    if column in JSON_COLUMNS:
        try:
            return json.loads(text) if text else []
        except json.JSONDecodeError:
            return []
    if column in LIST_COLUMNS:
        if not text:
            return []
        parts = [part.strip() for part in text.split(LIST_SEPARATOR.strip()) if part.strip()]
        if column in ID_LIST_COLUMNS:
            return [int(float(part)) for part in parts if part.replace(".", "").isdigit()]
        return parts
    if column in BOOL_COLUMNS:
        return text.lower() in {"yes", "true", "1"}
    if column in INT_COLUMNS:
        try:
            return int(float(text)) if text else None
        except ValueError:
            return None
    return text or None


def write_workbook(archive) -> bytes:
    book = Workbook()
    book.remove(book.active)

    meta = book.create_sheet(META_SHEET)
    meta.append(["key", "value"])
    meta.append(["version", archive["version"]])
    meta.append(["username", archive["username"]])
    for cell in meta[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT

    for name, columns in SHEETS.items():
        sheet = book.create_sheet(name)
        sheet.append(columns)
        for cell in sheet[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        sheet.freeze_panes = "A2"

        rows = archive.get(name, [])
        for row in rows:
            sheet.append([_to_cell(column, row.get(column)) for column in columns])

        # Wide enough to read without being unusable — free text is capped.
        for index, column in enumerate(columns, start=1):
            widths = [len(column)] + [
                len(str(_to_cell(column, row.get(column)))) for row in rows
            ]
            sheet.column_dimensions[get_column_letter(index)].width = min(
                max(max(widths) + 2, 10), 60
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def read_workbook(handle):
    """Parse a workbook this app wrote back into the archive shape."""
    try:
        book = load_workbook(handle, read_only=True, data_only=True)
    except Exception:
        raise serializers.ValidationError(
            {"file": "That file isn’t a readable spreadsheet."}
        )

    archive = {"version": ARCHIVE_VERSION, "username": ""}

    if META_SHEET in book.sheetnames:
        for key, value in book[META_SHEET].iter_rows(min_row=2, values_only=True):
            if key == "version" and value is not None:
                archive["version"] = int(value)
            elif key == "username" and value is not None:
                archive["username"] = str(value)

    for name, columns in SHEETS.items():
        rows = []
        if name in book.sheetnames:
            header = None
            for values in book[name].iter_rows(values_only=True):
                if header is None:
                    header = [str(v).strip() if v is not None else "" for v in values]
                    continue
                if all(v is None or str(v).strip() == "" for v in values):
                    continue
                raw = dict(zip(header, values))
                rows.append(
                    {column: _from_cell(column, raw.get(column)) for column in columns}
                )
        archive[name] = rows

    book.close()
    return archive
